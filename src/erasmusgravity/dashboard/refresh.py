"""Refresh descriptive exploration without changing the published research sample."""
from __future__ import annotations

from collections import Counter
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import unicodedata

import geonamescache
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from ..cities import aggregate_country_to_org_flows, aggregate_org_to_country_flows, aggregate_org_to_org_flows
from ..grants import join_sender_grants, load_he_grants
from ..indicators import fetch_country_indicators, indicator_coverage, join_indicators
from ..ka1 import KA1_COLUMNS, PARTICIPANTS_COL, aggregate_country_flows, filter_he_studies
from .payload import clean_records, country_records, json_safe, weighted_records

DOWNLOAD_PAGE = 'https://erasmus-plus.ec.europa.eu/resources-and-tools/statistics-and-factsheets/statistics/for-researchers'
SOURCE_URL = 'https://ec.europa.eu/assets/eac/erasmus-plus/statistics/mobility/Erasmus-KA1-Mobility-Data-{year}.xlsx'
RESEARCH_KEYS = ('modelCoefficients', 'regressionResults', 'panelYears', 'covidEventStudy',
                 'covidEventStudyEmployment', 'redditByCountry', 'redditMeta', 'sectorEarnings')


def read_snapshot(path: Path) -> dict:
    match = re.search(r'<script id="dashboard-data" type="application/json">(.*?)</script>', path.read_text(), re.S)
    if not match:
        raise ValueError('Snapshot has no dashboard-data payload')
    return json.loads(match[1])


def read_mobility(path: Path, year: int, cache: Path) -> tuple[pd.DataFrame, list[dict]]:
    """Stream the million-row Excel source; preserve participant weights and NA country code."""
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    manifest, csv = cache / 'filtered-source.json', cache / 'filtered.csv.gz'
    if manifest.exists() and csv.exists():
        saved = json.loads(manifest.read_text())
        if saved['sha256'] == digest and saved['year'] == year:
            frame = pd.read_csv(csv, keep_default_na=False)
            print(f'Reused {len(frame):,} filtered source rows', flush=True)
            return frame, saved['diagnostics']
    cache.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb[f'KA1 mobilities {year}']
    source = sheet.iter_rows(values_only=True)
    headers = next(source)
    indices = {name: headers.index(name) for name in KA1_COLUMNS}
    rows, counts, weights = [], Counter(), Counter()
    def record(stage, amount):
        counts[stage] += 1
        weights[stage] += amount
    for index, row in enumerate(source, 1):
        if index % 200000 == 0:
            print(f'Read {index:,} source rows', flush=True)
        try:
            amount = float(row[indices[PARTICIPANTS_COL]] or 0)
        except (ValueError, TypeError):
            raise ValueError(f'Invalid participant weight at source row {index + 1}')
        if amount < 0 or not amount.is_integer():
            raise ValueError(f'Invalid participant weight {amount} at row {index + 1}')
        amount = int(amount)
        record('raw', amount)
        if row[indices['Field']] != 'Higher Education':
            continue
        record('higher_education', amount)
        if str(row[indices['Participant Profile']]).lower() not in ('learner', 'learners'):
            continue
        record('learners_only', amount)
        activity = str(row[indices['Activity (mob)']])
        if not (activity.startswith('HE-SMS') or 'student mobility for studies' in activity.lower()):
            continue
        record('studies_only', amount)
        try:
            age = float(row[indices['Participant Age']])
        except (ValueError, TypeError):
            continue
        if age < 18 or not np.isfinite(age):
            continue
        record('age_18_plus', amount)
        if not str(row[indices['Mobility Start Month']]).startswith(f'{year}-'):
            raise ValueError('Mobility start year differs from the requested release')
        rows.append({name: row[column] for name, column in indices.items()})
    wb.close()
    filtered, extra = filter_he_studies(pd.DataFrame(rows, columns=KA1_COLUMNS))
    filtered = filtered.loc[filtered['sending_code'].ne(filtered['receiving_code']) & filtered[PARTICIPANTS_COL].gt(0)].copy()
    diagnostics = [{'step': key, 'rows': counts[key], 'participants': weights[key]} for key in counts]
    diagnostics += [extra.iloc[-1].to_dict(), {'step': 'cross_border_positive', 'rows': len(filtered), 'participants': int(filtered[PARTICIPANTS_COL].sum())}]
    # Empty endpoint fields stay explicit when read back; country code NA is never an NA token.
    filtered.to_csv(csv, index=False)
    manifest.write_text(json.dumps(json_safe({'year': year, 'sha256': digest, 'diagnostics': diagnostics}), indent=2))
    print(f'Filtered {len(filtered):,} rows / {int(filtered[PARTICIPANTS_COL].sum()):,} participants', flush=True)
    return filtered.fillna(''), diagnostics


def university_tables(filtered: pd.DataFrame) -> dict[str, list[dict]]:
    frame = filtered.copy()
    for column in ('Sending City', 'Receiving City'):
        frame[column] = frame[column].replace({'': 'Not recorded', '-': 'Not recorded'}).fillna('Not recorded')
    for column in ('Sending Organization', 'Receiving Organization'):
        frame[column] = frame[column].replace({'': None, '-': None})
    levels = {'countryUni': aggregate_country_to_org_flows(frame),
              'uniCountry': aggregate_org_to_country_flows(frame),
              'uniUni': aggregate_org_to_org_flows(frame)}
    expected = int(filtered[PARTICIPANTS_COL].sum())
    for name, table in levels.items():
        if int(table.flow_count.sum()) != expected:
            raise ValueError(f'{name} lost or duplicated participants')
    return {key: json_safe(value.to_dict('records')) for key, value in levels.items()}


def _normalize(value) -> str:
    text = unicodedata.normalize('NFKD', str(value or ''))
    return re.sub(r'[^a-z0-9]', '', ''.join(c for c in text if not unicodedata.combining(c)).lower())


def map_coordinates(snapshot: dict, tables: dict) -> list[dict]:
    """Reuse year-independent geography, then fill city matches from GeoNames."""
    lookup = {}
    names = {_normalize(r['country']): r['code'] for r in snapshot['countries']}
    for item in snapshot.get('mapCoordinates', []):
        lookup[(item['code'], _normalize(item['city']))] = item
    for point in snapshot.get('orgFlowLevels', {}).get('countryUni', {}).get('all', []):
        code = names.get(_normalize(point['n']))
        if code and point['m'] == 'city':
            lookup.setdefault((code, _normalize(point['c'])), {'code': code, 'city': point['c'], 'lat': point['la'], 'lon': point['lo'], 'level': 'city'})
    country_points = {}
    for item in lookup.values():
        if item['level'] == 'country':
            country_points[item['code']] = item
    for point in snapshot.get('orgFlowLevels', {}).get('uniCountry', {}).get('all', []):
        if point.get('dc'):
            code = point['dc']
            country_points[code] = {'code': code, 'city': '', 'lat': point['la'], 'lon': point['lo'], 'level': 'country'}
    desired = set()
    for row in tables['uniUni']:
        desired.add((row['sending_code'], _normalize(row['sending_city'])))
        desired.add((row['receiving_code'], _normalize(row['receiving_city'])))
    # Population order resolves same-name cities within a country consistently.
    cities = sorted(geonamescache.GeonamesCache().get_cities().values(), key=lambda r: r['population'], reverse=True)
    for city in cities:
        code = {'GR': 'EL', 'GB': 'UK'}.get(city['countrycode'], city['countrycode'])
        for name in [city['name'], *city.get('alternatenames', [])]:
            key = (code, _normalize(name))
            if key in desired and key not in lookup:
                lookup[key] = {'code': code, 'city': name, 'lat': round(city['latitude'], 3), 'lon': round(city['longitude'], 3), 'level': 'city'}
    return [lookup[key] for key in sorted(desired) if key in lookup] + list(country_points.values())


def extend_history(snapshot: dict, filtered: pd.DataFrame, year: int) -> tuple[list[dict], list[str]]:
    """Append national counts within the archived panel's country coverage; do not refit models."""
    original = [r for r in snapshot['timeSeries'] if r['y'] != year]
    supported = {r[key] for r in original for key in ('s', 'd')}
    history = filtered.loc[filtered.sending_code.isin(supported) & filtered.receiving_code.isin(supported)
                           & pd.to_numeric(filtered.participant_age).le(100)]
    flows = aggregate_country_flows(history)
    extension = [{'s': r.sending_code, 'd': r.receiving_code, 'dn': r.receiving_country,
                  'y': year, 'f': int(r.flow_count)} for r in flows.itertuples()]
    return original + extension, sorted(supported)


def refresh_snapshot(snapshot: dict, filtered: pd.DataFrame, indicators: pd.DataFrame,
                     grants: pd.DataFrame, year: int) -> dict:
    updated = deepcopy(snapshot)
    flows = aggregate_country_flows(filtered)
    joined = join_sender_grants(join_indicators(flows, indicators), grants)
    if len(joined) != len(flows) or joined.flow_count.sum() != flows.flow_count.sum():
        raise ValueError('Economic joins changed the mobility counts')
    tables = university_tables(filtered)
    updated.update(flows=clean_records(joined), countries=country_records(joined), weighted=weighted_records(joined),
                   orgFlowTables=tables, mapCoordinates=map_coordinates(snapshot, tables))
    # The current UI uses explicit coordinate references. No old mobility counts remain in map layers.
    updated['orgFlowLevels'] = {key: {'points': [], 'all': []} for key in tables}
    updated['timeSeries'], supported = extend_history(snapshot, filtered, year)
    archive_end = max(r['year'] for r in snapshot['panelYears'])
    updated['explorationMeta'] = {
        'year': year, 'indicatorYear': year,
        'grantCallYear': int(grants['grant_call_year'].max()) if 'grant_call_year' in grants else 2022,
        'researchYear': 2022,
        'researchPanelEnd': archive_end,
        'historyStart': min(r['y'] for r in updated['timeSeries']), 'historyEnd': year,
        'historyBreakYear': snapshot.get('explorationMeta', {}).get('historyBreakYear', year),
        'historyCountryCodes': supported,
        'historyNote': f'2014–{archive_end}: geolocated archive. {year}: official KA1 release, restricted to the archive’s country coverage and ages 18–100. The separate latest point marks the source change; coverage and revisions can affect comparisons. Country profiles use the full latest release.',
        'participants': int(flows.flow_count.sum()), 'countryPairs': len(flows),
        'universityRows': len(tables['uniUni']), 'universityParticipants': sum(r['flow_count'] for r in tables['uniUni']),
        'scope': 'Higher Education learners, study mobility, age 18+, cross-border, positive participant weights; calendar year of mobility start.',
        'sourceUrl': SOURCE_URL.format(year=year), 'downloadPage': DOWNLOAD_PAGE,
        'indicatorSources': {
            'gdp': f'https://api.worldbank.org/v2/country/all/indicator/NY.GDP.PCAP.CD?format=json&date={year}&per_page=20000',
            'prices': f'https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/prc_ppp_ind?time={year}&na_item=PLI_EU27_2020&ppp_cat=A01',
            'earnings': f'https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/earn_nt_net?time={year}&currency=EUR&estruct=NET&ecase=P1_NCH_AW100',
        },
        'indicatorCoverage': indicator_coverage(joined).to_dict('records'),
    }
    for key in RESEARCH_KEYS:
        if updated[key] != snapshot[key]:
            raise AssertionError(f'Research/reference payload changed: {key}')
    return json_safe(updated)


def build_refresh(snapshot: Path, source: Path, year: int, cache: Path, annex: Path) -> dict:
    original = read_snapshot(snapshot)
    filtered, diagnostics = read_mobility(source, year, cache)
    indicator_path = cache / f'country-indicators-{year}.json'
    if indicator_path.exists():
        indicators = pd.DataFrame(json.loads(indicator_path.read_text()))
    else:
        indicators = fetch_country_indicators(year)
        indicator_path.write_text(json.dumps(json_safe(indicators.to_dict('records')), indent=2))
    grant_releases = pd.read_excel(annex, sheet_name='KA1_estimated paxs', usecols=['Field', 'Call Year'])
    available = grant_releases.loc[grant_releases.Field.eq('Higher Education'), 'Call Year']
    grant_year = int(available.loc[available.le(year)].max())
    grants = load_he_grants(annex, call_year=grant_year)
    if grants.empty:
        raise ValueError('Annex has no compatible grant/participant records')
    result = refresh_snapshot(original, filtered, indicators, grants, year)
    result['explorationMeta'].update(sourceFile=source.name, sourceSha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                                   checkedAt=datetime.now(timezone.utc).date().isoformat(), diagnostics=diagnostics,
                                   grantSourceFile=annex.name, grantSourceSha256=hashlib.sha256(annex.read_bytes()).hexdigest())
    return result
